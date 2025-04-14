import streamlit as st
from docx import Document
from io import BytesIO
import pypandoc  # Optional for PDF

#extracting template from the uploaded file

def extract_template_format(docx_file):
    doc = Document(docx_file)
    header = doc.sections[0].header.paragraphs[0].text
    footer = doc.sections[0].footer.paragraphs[0].text
    return header, footer

#applying format on the raw file

def apply_format(target_file, header, footer):
    doc = Document(target_file)
    doc.sections[0].header.paragraphs[0].text = header
    doc.sections[0].footer.paragraphs[0].text = footer

    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output

#converting the file into pdf

def convert_to_pdf(docx_bytes):
    with open("temp.docx", "wb") as f:
        f.write(docx_bytes.read())
    output_pdf = pypandoc.convert_file("temp.docx", "pdf", outputfile="output.pdf")
    with open("output.pdf", "rb") as f:
        return f.read()
    

#developing UI using streamlit

st.title("📄 Document Formatter")

template_docx = st.file_uploader("Upload Template Document", type=["docx"])
target_docs = st.file_uploader("Upload Target Document(s)", type=["docx"], accept_multiple_files=True)

format_option = st.selectbox("Output Format", ["DOCX", "PDF"])

if st.button("Apply Format") and template_docx and target_docs:
    header, footer = extract_template_format(template_docx)
    for file in target_docs:
        formatted_doc = apply_format(file, header, footer)

        filename = file.name.replace(".docx", f"_formatted.{format_option.lower()}")

        if format_option == "PDF":
            pdf_bytes = convert_to_pdf(formatted_doc)
            st.download_button("Download " + filename, data=pdf_bytes, file_name=filename, mime="application/pdf")
        else:
            st.download_button("Download " + filename, data=formatted_doc, file_name=filename, mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
